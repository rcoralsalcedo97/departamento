"""Second-pass rules: budget classes, integrity guard, foreign-tenant evidence, Apify auth and input mapping."""
import json
from pathlib import Path

import pandas as pd
import pytest

from src.config import ROOT
from src.http_client import NetworkBlocked
from src.normalize.foreign_tenant import assess_foreign_tenant
from src.normalize.normalize import normalize_listing
from src.scoring.scoring import budget_class, classify, integrity_issues, rank, score_all
from src.sources.apify_client import build_actor_input, resolve_apify_auth
from src.sources.navent import map_navent_record

BASE = dict(source="urbania", source_listing_id="1", source_url="https://urbania.pe/inmueble/x-1", operation="Alquiler",
            property_type="Departamento", bedrooms=2, rent_usd=900.0, rent_usd_basis="PUBLISHED", district="Miraflores",
            inside_district_polygon=True, coord_precision="EXACT", latitude=-12.12, longitude=-77.03,
            built_area_m2=75.0, active_status="LIKELY_ACTIVE", dist_major_road_m=400, major_road_name="Av. X",
            dist_arterial_road_m=300, dist_nightclub_m=900, bars_within_m=0, restaurants_within_m=1,
            is_canonical=True, dist_outside_district_m=0, title="Departamento 2 dormitorios")


def test_budget_classes(cfg):
    assert budget_class({**BASE, "estimated_total_monthly_usd": 980.0}, cfg)[0] == "STRICT_ALL_IN"
    assert budget_class({**BASE, "estimated_total_monthly_usd": 1040.0}, cfg)[0] == "BASE_RENT_COMPLIANT"
    assert budget_class({**BASE, "estimated_total_monthly_usd": None}, cfg)[0] == "BASE_RENT_COMPLIANT"
    assert budget_class({**BASE, "maintenance_included_in_rent": True}, cfg)[0] == "STRICT_ALL_IN"
    assert budget_class({**BASE, "rent_usd": 1080.0}, cfg)[0] == "STRETCH"


def test_strict_all_in_preferred_unless_clearly_better(cfg):
    strict = {**BASE, "source_listing_id": "S", "source_url": "https://urbania.pe/inmueble/s", "estimated_total_monthly_usd": 990.0,
              "maintenance_usd": 90.0}
    unknown = {**BASE, "source_listing_id": "U", "source_url": "https://urbania.pe/inmueble/u", "furnished": True,
               "estimated_total_monthly_usd": None}
    out = rank(score_all(pd.DataFrame([unknown, strict]), cfg), cfg["budget"]["strict_preference_margin"])
    assert list(out["source_listing_id"]) == ["S", "U"]


def test_low_confidence_noise_is_not_demoted(cfg):
    loud_uncertain = {**BASE, "source_listing_id": "L", "source_url": "https://urbania.pe/inmueble/l",
                      "dist_major_road_m": 20, "dist_nightclub_m": 90, "bars_within_m": 5,
                      "coord_precision": "APPROXIMATE", "built_area_m2": 120.0, "furnished": True,
                      "estimated_total_monthly_usd": 950.0, "maintenance_usd": 50.0}
    plain = {**BASE, "source_listing_id": "P", "source_url": "https://urbania.pe/inmueble/p", "built_area_m2": 50.0,
             "estimated_total_monthly_usd": 990.0, "maintenance_usd": 90.0}
    out = rank(score_all(pd.DataFrame([plain, loud_uncertain]), cfg), cfg["budget"]["strict_preference_margin"])
    row = out[out["source_listing_id"] == "L"].iloc[0]
    assert row["noise_confidence"] == "LOW"
    assert row["rank_in_category"] == 1       # not pushed behind quieter units on uncertain geodata


def test_integrity_guard_rejects_every_synthetic_fixture(cfg, fx):
    raw = json.loads((ROOT / "tests" / "fixtures" / "synthetic_navent.json").read_text(encoding="utf-8"))
    prod = {**cfg, "_production": True}
    for rec in raw:
        lst = normalize_listing(map_navent_record(dict(rec), rec["_source"], "https://example.com"), cfg, fx)
        cat, reasons, _ = classify(lst.model_dump(), prod)
        assert cat == "EXCLUDED" and any(r.startswith("INTEGRITY") for r in reasons), lst.source_url


def test_integrity_guard_keeps_real_spanish_text(cfg):
    ok = {**BASE, "description": "Cerca de todo, por ejemplo Larcomar y el parque Kennedy."}
    assert integrity_issues(ok) == []
    assert classify(ok, {**cfg, "_production": True})[0] == "PRIMARY"


def test_production_outputs_contain_no_demo_data():
    """Scans the real deliverables when they exist; the pipeline's Gate 10 runs the same scan."""
    from openpyxl import load_workbook
    from src.qa.final_qa import DEMO_RX
    xlsx = ROOT / "outputs" / "Miraflores_Rental_Shortlist_REAL.xlsx"
    if not xlsx.exists():
        pytest.skip("no production run yet")
    wb = load_workbook(xlsx)
    hits = [c.coordinate for ws in wb for row in ws.iter_rows() for c in row
            if isinstance(c.value, str) and DEMO_RX.search(c.value)]
    assert hits == []
    txt = (ROOT / "outputs" / "Contact_Templates_REAL.txt").read_text(encoding="utf-8")
    assert not DEMO_RX.search(txt)
    import pymupdf
    pdf_text = "".join(p.get_text() for p in pymupdf.open(ROOT / "outputs" / "Miraflores_Rental_Executive_Report_REAL.pdf"))
    assert not DEMO_RX.search(pdf_text)


def test_foreign_tenant_is_evidence_based():
    assert assess_foreign_tenant("Depa", "Lindo departamento en Miraflores.", None, None)[0] == "UNKNOWN"
    assert assess_foreign_tenant("Depa", "Se aceptan extranjeros con pasaporte.", None, None)[0] == "HIGH"
    assert assess_foreign_tenant("Depa", "Requisito: aval con inmueble.", None, None)[0] == "POTENTIAL_FRICTION"
    assert assess_foreign_tenant("Depa", "Sin aval. Alquiler temporal.", None, None)[0] == "MEDIUM"


def test_actor_input_mapping_to_schema():
    props = {"startUrls": {"type": "array", "editor": "requestListSources"},
             "priceMax": {"type": "integer"}, "currency": {"type": "string", "enum": ["PEN", "USD"]},
             "operationType": {"type": "string", "enum": ["venta", "alquiler"]},
             "maxItems": {"type": "integer"}, "includeDetails": {"type": "boolean"}}
    out, notes = build_actor_input(props, {"start_urls": ["https://urbania.pe/buscar/x"], "max_price": 1150,
                                           "currency": "USD", "operation": "rent", "with_details": True,
                                           "min_bedrooms": 1}, 15)
    assert out == {"startUrls": [{"url": "https://urbania.pe/buscar/x"}], "priceMax": 1150, "currency": "USD",
                   "operationType": "alquiler", "includeDetails": True, "maxItems": 15}
    assert any("min_bedrooms" in n for n in notes)


class _Blocked:
    def request_json(self, *a, **k):
        raise NetworkBlocked("api.apify.com: 403")


def test_auth_reports_blocked_network_without_secret(monkeypatch):
    dummy = "dummy-" + "credential"          # deliberately not token-shaped (keeps the repo leak scan clean)
    monkeypatch.setenv("APIFY_TOKEN", dummy)
    auth = resolve_apify_auth(_Blocked())
    shown = "\n".join(auth.display())
    assert not auth.available and auth.method == "ENVIRONMENT_VARIABLE"
    assert dummy not in shown and "unreachable" in shown
