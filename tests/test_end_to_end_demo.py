"""Offline end-to-end run on synthetic fixtures (no network): exercises every gate and all deliverables."""
from openpyxl import load_workbook

from src import pipeline


def test_demo_pipeline_builds_all_deliverables(tmp_path):
    assert pipeline.main(["--mode", "demo", "--out", str(tmp_path)]) == 0
    out = tmp_path
    wb = load_workbook(out / "PREVIEW_SYNTHETIC_Miraflores_Rental_Shortlist_EN.xlsx")
    assert wb.sheetnames == ["CLIENT_TOP_PICKS", "EXECUTIVE_SHORTLIST", "ALL_MATCHES", "STRETCH_NEGOTIABLE", "NEAR_MISSES",
                             "SOURCE_AUDIT", "METHODOLOGY", "CONTACT_GUIDE"]
    assert len([r for r in wb["CLIENT_TOP_PICKS"].iter_rows(min_row=5) if r[0].value]) == 10
    assert wb["CLIENT_TOP_PICKS"].max_column <= 14
    ws = wb["EXECUTIVE_SHORTLIST"]
    assert "SYNTHETIC" in str(ws["A3"].value)
    assert any(c.hyperlink for row in ws.iter_rows() for c in row)
    assert (out / "PREVIEW_SYNTHETIC_Miraflores_Rental_Executive_Report_EN.pdf").stat().st_size > 20000
    assert "¿El departamento sigue disponible?" in (out / "PREVIEW_SYNTHETIC_Contact_Templates_EN.txt").read_text(encoding="utf-8")
    # Hindi copies: same structure, first sheet is the client Top 10, same Spanish message
    hi = load_workbook(out / "PREVIEW_SYNTHETIC_Miraflores_Rental_Shortlist_HI.xlsx")
    assert hi.sheetnames[0] == "शीर्ष विकल्प" and len(hi.sheetnames) == len(wb.sheetnames)
    assert (out / "PREVIEW_SYNTHETIC_Miraflores_Rental_Executive_Report_HI.pdf").stat().st_size > 20000
    hi_txt = (out / "PREVIEW_SYNTHETIC_Contact_Templates_HI.txt").read_text(encoding="utf-8")
    assert "¿El departamento sigue disponible?" in hi_txt and "क्या अपार्टमेंट अभी उपलब्ध है?" in hi_txt
