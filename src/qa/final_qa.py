"""Gate 10 — final deliverable QA (automated; results go to RUN_LOG.md).

Checks the finished files, not just the data: the workbook opens with the right sheets and
live hyperlinks; every PDF page renders with no text running off the page or overlapping;
the Top-10 links answer; no demo/synthetic data or credential leaked; and the Top 10 obey
the hard requirements. Page renders are written to data/processed/qa/ for a human look.
"""
from __future__ import annotations

import math
import re
from pathlib import Path

import pandas as pd

from .. import config as K

DEMO_RX = re.compile(r"\(DEMO|\bDEMO\b|Calle Ejemplo|Av\. Ejemplo|SYNTHETIC|Synthetic|synthetic|example\.com|\bfake listing")
TOKEN_RX = re.compile(r"apify_api_[A-Za-z0-9]{16,}")
REQUIRED_SHEETS = ["CLIENT_TOP_PICKS", "EXECUTIVE_SHORTLIST", "ALL_MATCHES", "STRETCH_NEGOTIABLE", "NEAR_MISSES",
                   "SOURCE_AUDIT", "METHODOLOGY", "CONTACT_GUIDE"]


def _sheet_names() -> dict[str, dict[str, str]]:
    from ..reporting.excel_hi import SHEETS
    return {"EN": {s: s for s in REQUIRED_SHEETS}, "HI": {s: SHEETS[s] for s in REQUIRED_SHEETS}}


class _Names(dict):
    def __missing__(self, key):
        self.update(_sheet_names())
        return dict.__getitem__(self, key)


SHEET_NAMES = _Names()
DEVANAGARI = re.compile(r"[\u0900-\u097F]")


def _num(v):
    try:
        f = float(v)
        return None if math.isnan(f) else f
    except (TypeError, ValueError):
        return None


class Report:
    def __init__(self):
        self.lines: list[str] = []
        self.failed = False

    def add(self, status: str, label: str, detail: str = "") -> None:
        if status == "FAIL":
            self.failed = True
        self.lines.append(f"{status:4}  {label}" + (f" — {detail}" if detail else ""))


# --------------------------------------------------------------------------- workbook
def xlsx_checks(path: Path, rep: Report, production: bool, tag: str = "EN") -> list[str]:
    from openpyxl import load_workbook
    try:
        wb = load_workbook(path)
    except Exception as exc:  # noqa: BLE001
        rep.add("FAIL", f"[{tag}] 1. XLSX opens", str(exc))
        return []
    names = SHEET_NAMES[tag]
    missing = [names[s] for s in REQUIRED_SHEETS if names[s] not in wb.sheetnames]
    first_ok = wb.sheetnames[0] == names["CLIENT_TOP_PICKS"]
    links = sum(1 for ws in wb for row in ws.iter_rows() for c in row if c.hyperlink)
    hidden = [f"{ws.title}!{k}" for ws in wb for k, d in ws.column_dimensions.items() if d.hidden] + \
        [ws.title for ws in wb if ws.sheet_state != "visible"]
    rep.add("PASS" if not hidden else "FAIL", f"[{tag}] 1c. No hidden columns or sheets", ", ".join(hidden[:5]))
    rep.add("PASS" if not missing and first_ok else "FAIL", f"[{tag}] 1. XLSX opens with the expected sheets",
            f"{len(wb.sheetnames)} sheets, first = {wb.sheetnames[0]}, {links} hyperlinks"
            + (f", missing {missing}" if missing else ""))
    clipped = []
    for name in (names["CLIENT_TOP_PICKS"], names["EXECUTIVE_SHORTLIST"]):
        if name not in wb.sheetnames:
            continue
        ws = wb[name]
        for row in ws.iter_rows(min_row=5):
            for c in row:
                if c.value is None or c.hyperlink or (c.alignment and c.alignment.wrap_text):
                    continue
                width = ws.column_dimensions[c.column_letter].width or 8.43
                if len(str(c.value)) > width * 1.15:
                    clipped.append(f"{name}!{c.coordinate}")
    rep.add("PASS" if not clipped else "FAIL", f"[{tag}] 4. No clipped (unwrapped, overflowing) cells",
            ", ".join(clipped[:8]) + (" …" if len(clipped) > 8 else ""))
    urls = []
    if names["CLIENT_TOP_PICKS"] in wb.sheetnames:
        for row in wb[names["CLIENT_TOP_PICKS"]].iter_rows(min_row=5):
            for c in row:
                if c.hyperlink and c.hyperlink.target:
                    urls.append(c.hyperlink.target)
    if production:
        hits = [f"{ws.title}!{c.coordinate}" for ws in wb for row in ws.iter_rows() for c in row
                if isinstance(c.value, str) and DEMO_RX.search(c.value)]
        rep.add("PASS" if not hits else "FAIL", f"[{tag}] 6. No demo/synthetic data in the workbook", ", ".join(hits[:5]))
    return urls


def render_sheet_png(path: Path, sheet: str, out_png: Path) -> None:
    """Approximate on-screen rendering of a sheet (column widths, wrapping) for visual review."""
    import html as H
    from openpyxl import load_workbook
    from playwright.sync_api import sync_playwright
    ws = load_workbook(path)[sheet]
    widths = [(ws.column_dimensions[c.column_letter].width or 8.43) * 7.2 for c in ws[4]]
    rows_html = []
    for r in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 16)):
        cells = []
        for j, c in enumerate(r):
            v = "" if c.value is None else str(c.value)
            style = f"width:{widths[j] if j < len(widths) else 60:.0f}px;max-width:{widths[j] if j < len(widths) else 60:.0f}px;"
            wrap = c.alignment is not None and c.alignment.wrap_text
            style += "white-space:normal;" if wrap else "white-space:nowrap;overflow:hidden;"
            if c.hyperlink:
                v = f"<u style='color:#1f5fbf'>{H.escape(v)}</u>"
            else:
                v = H.escape(v)
            bg = c.fill.fgColor.rgb if c.fill and c.fill.fill_type == "solid" else None
            if isinstance(bg, str) and len(bg) == 8:
                style += f"background:#{bg[2:]};"
            if c.font and c.font.color and isinstance(c.font.color.rgb, str) and len(c.font.color.rgb) == 8:
                style += f"color:#{c.font.color.rgb[2:]};"
            if c.font and c.font.bold:
                style += "font-weight:bold;"
            cells.append(f"<td style='{style}'>{v}</td>")
        rows_html.append("<tr>" + "".join(cells) + "</tr>")
    doc = ("<html><body style='margin:0;font:12px Calibri,Liberation Sans,Arial'><table style='border-collapse:"
           "collapse;table-layout:fixed' border=1 cellpadding=3>" + "".join(rows_html) + "</table></body></html>")
    tmp = out_png.with_suffix(".html")
    tmp.write_text(doc, encoding="utf-8")
    exe = next(iter(sorted(Path("/opt/pw-browsers").glob("chromium-*/chrome-linux/chrome"))), None)
    with sync_playwright() as p:
        b = p.chromium.launch(executable_path=str(exe)) if exe else p.chromium.launch()
        pg = b.new_page(viewport={"width": 1366, "height": 768})
        pg.goto(tmp.resolve().as_uri())
        pg.screenshot(path=str(out_png), full_page=False)
        b.close()


# --------------------------------------------------------------------------- PDF
def pdf_checks(path: Path, rep: Report, out_dir: Path, production: bool, tag: str = "EN") -> list[str]:
    import pymupdf
    try:
        doc = pymupdf.open(path)
    except Exception as exc:  # noqa: BLE001
        rep.add("FAIL", f"[{tag}] 2. PDF opens", str(exc))
        return []
    n = len(doc)
    overflow, overlaps, links, text_all = [], [], [], []
    out_dir.mkdir(parents=True, exist_ok=True)
    for i, page in enumerate(doc, start=1):
        page.get_pixmap(dpi=80).save(out_dir / f"pdf_page_{i:02d}.png")
        w = page.rect.width
        lines = []
        for block in page.get_text("dict")["blocks"]:
            for line in block.get("lines", []):
                bbox = pymupdf.Rect(line["bbox"])
                txt = "".join(sp["text"] for sp in line["spans"]).strip()
                if not txt:
                    continue
                text_all.append(txt)
                if bbox.x1 > w - 8 or bbox.x0 < 8:
                    overflow.append(f"p{i}: '{txt[:30]}'")
                lines.append((bbox, txt))
        for a in range(len(lines)):
            for b in range(a + 1, len(lines)):
                inter = lines[a][0] & lines[b][0]
                if inter.is_empty:
                    continue
                small = min(lines[a][0].get_area(), lines[b][0].get_area()) or 1
                if inter.get_area() / small > 0.35:
                    overlaps.append(f"p{i}: '{lines[a][1][:20]}' × '{lines[b][1][:20]}'")
        links += [ln.get("uri") for ln in page.get_links() if ln.get("uri")]
    rep.add("PASS" if 6 <= n <= 12 else "WARN", f"[{tag}] 2. PDF opens and every page renders",
            f"{n} pages (target 6–12); page images in {out_dir}")
    rep.add("PASS" if not overflow else "FAIL", f"[{tag}] 4b. No text running off the page", "; ".join(overflow[:5]))
    rep.add("PASS" if not overlaps else "FAIL", f"[{tag}] 5. No overlapping text", "; ".join(overlaps[:5]))
    if production:
        joined = "\n".join(text_all)
        hits = DEMO_RX.findall(joined)
        rep.add("PASS" if not hits else "FAIL", f"[{tag}] 6b. No demo/synthetic wording in the PDF", ", ".join(sorted(set(hits))))
    return links


# --------------------------------------------------------------------------- links, leaks, rules
def link_checks(urls: list[str], http, rep: Report, label: str) -> None:
    if http is None:
        rep.add("SKIP", label, "no network client")
        return
    listing = [u for u in dict.fromkeys(urls) if not re.search(r"wa\.me|google\.com/maps|openstreetmap", u)]
    other = [u for u in dict.fromkeys(urls) if u not in listing]
    dead, protected, alive = [], [], 0
    for u in listing:
        try:
            r = http.request_json("GET", u, timeout=25)
            if r.status_code in (404, 410):
                dead.append(u)
            elif r.status_code in (401, 403, 429, 503):
                protected.append(u)
            else:
                alive += 1
        except Exception as exc:  # noqa: BLE001
            protected.append(f"{u} ({type(exc).__name__})")
    malformed = [u for u in other if not u.startswith("https://")]
    status = "FAIL" if dead or malformed else ("WARN" if protected else "PASS")
    rep.add(status, label, f"{alive} listing links answer, {len(protected)} bot-protected/unreachable to scripts "
                           f"(open manually), {len(dead)} dead; {len(other)} map/WhatsApp links well-formed"
            + (f"; dead: {dead[:3]}" if dead else ""))


def leak_checks(paths: list[Path], token: str | None, rep: Report) -> None:
    import subprocess
    scan = list(paths)
    for pattern in ("*.md", "config/*", "outputs/*", "data/processed/*.csv", "data/processed/*.json"):
        scan += [p for p in K.ROOT.glob(pattern) if p.is_file()]
    try:
        tracked = subprocess.run(["git", "ls-files"], cwd=K.ROOT, capture_output=True, text=True).stdout.split()
        scan += [K.ROOT / t for t in tracked]
    except Exception:  # noqa: BLE001
        pass
    leaked = []
    for p in dict.fromkeys(scan):
        if not p.is_file() or p.name == ".env" or p.stat().st_size > 30_000_000:
            continue
        try:
            data = p.read_bytes()
        except OSError:
            continue
        if TOKEN_RX.search(data.decode("latin-1")) or (token and token.encode() in data):
            leaked.append(str(p.relative_to(K.ROOT)))
    rep.add("PASS" if not leaked else "FAIL", "7. No secret/token in any file", ", ".join(leaked))


def rule_checks(ranked: pd.DataFrame, cfg: dict, rep: Report, production: bool) -> None:
    top = ranked[ranked["category"] == "PRIMARY"].sort_values("rank_in_category").head(cfg["shortlist"]["top_n"])
    district = cfg["location"]["district"]
    bad_d = [r["source_url"] for r in top.to_dict("records")
             if not (r.get("district") == district and r.get("inside_district_polygon") is not False)
             and r.get("inside_district_polygon") is not True]
    rep.add("PASS" if not bad_d and len(top) else "FAIL", "8. Every Top-10 listing is in Miraflores",
            f"{len(top)} checked" + (f"; problems: {bad_d[:3]}" if bad_d else ""))
    rep.add("PASS" if top["bedrooms"].isin([1, 2]).all() else "FAIL", "9. Every Top-10 listing has 1–2 bedrooms")
    wrong = []
    for r in top.to_dict("records"):
        rent, total = _num(r.get("rent_usd")), _num(r.get("estimated_total_monthly_usd"))
        bc = r.get("budget_class")
        if rent is None or rent > cfg["budget"]["target_max_rent"]:
            wrong.append(f"{r['source_url']} rent {rent}")
        elif bc == "STRICT_ALL_IN" and not (total is not None and total <= cfg["budget"]["target_max_total"]
                                            or r.get("maintenance_included_in_rent") is True):
            wrong.append(f"{r['source_url']} marked STRICT without known total")
        elif bc == "BASE_RENT_COMPLIANT" and total is not None and total <= cfg["budget"]["target_max_total"]:
            wrong.append(f"{r['source_url']} total ≤ budget but not STRICT")
    rep.add("PASS" if not wrong else "FAIL", "10. Budget classification consistent", "; ".join(wrong[:3]))
    groups = ranked["duplicate_group_id"]
    rep.add("PASS" if groups.notna().all() and groups.is_unique else "FAIL", "11. Duplicate groups",
            f"{groups.nunique()} canonical groups; {int((ranked['number_of_records'] > 1).sum())} merge ≥2 records")
    conf_ok = top["noise_confidence"].isin(["HIGH", "MEDIUM", "LOW"]).all() and \
        top["quietness_reason"].astype(str).str.contains("confidence").all()
    rep.add("PASS" if conf_ok else "FAIL", "12. Quietness claims state their confidence")
    if production:
        demo = [r["source_url"] for r in ranked.to_dict("records")
                if DEMO_RX.search(" ".join(str(r.get(k) or "") for k in ("title", "address", "source_url",
                                                                        "agency_name")))]
        rep.add("PASS" if not demo else "FAIL", "6c. Zero demo/synthetic records in the production dataset",
                f"{len(ranked)} records scanned" + (f"; found {demo[:3]}" if demo else ""))
    counts = top["active_status"].value_counts().to_dict()
    rep.add("INFO", "Top-10 availability", ", ".join(f"{k} {v}" for k, v in counts.items()))


def devanagari_checks(pdf: Path, rep: Report) -> None:
    """Hindi PDF: Devanagari is present, set in a Devanagari font, with no tofu / replacement / dotted-circle glyphs
    (Chromium draws U+25CC when a vowel sign cannot attach — the typical sign of broken shaping)."""
    import pymupdf
    doc = pymupdf.open(pdf)
    wrong_font, bad, n_dev = set(), [], 0
    for i, page in enumerate(doc, start=1):
        for block in page.get_text("dict")["blocks"]:
            for line in block.get("lines", []):
                for sp in line["spans"]:
                    txt = sp["text"]
                    if DEVANAGARI.search(txt):
                        n_dev += len(DEVANAGARI.findall(txt))
                        if "devanagari" not in sp["font"].lower():
                            wrong_font.add(sp["font"])
                    if any(ch in txt for ch in ("\ufffd", "\u25a1", "\u25cc", "\u0000")):
                        bad.append(f"p{i}: '{txt[:25]}'")
    fonts = {f[3] for page in doc for f in page.get_fonts()}
    # English prose leaking into the Hindi report (names, addresses, URLs and codes are allowed; sentences are not)
    prose = re.compile(r"\b(the|from|with|within|listing|says|approximate|mapped|major road|and|of|is|not|only|"
                       r"confirm|maintenance|rent|furnished|noise|bedroom)\b", re.I)
    leaks = []
    for i, page in enumerate(doc, start=1):
        for line in page.get_text().splitlines():
            clean = re.sub(r"https?://\S+|\S+\.xlsx|Miraflores Rental Shortlist", "", line)
            if prose.search(clean):
                leaks.append(f"p{i}: '{line.strip()[:60]}'")
    rep.add("PASS" if not leaks else "FAIL", "[HI] 19. No English prose left in the Hindi report",
            f"{len(leaks)} lines" + (f": {leaks[:4]}" if leaks else ""))
    rep.add("PASS" if n_dev > 2000 and not wrong_font and not bad else "FAIL",
            "[HI] 14. Devanagari renders with a Devanagari font (no boxes, broken signs or replacement glyphs)",
            f"{n_dev} Devanagari characters; fonts: {', '.join(sorted(x for x in fonts if 'Devanagari' in x))}"
            + (f"; non-Devanagari font used for Hindi: {sorted(wrong_font)}" if wrong_font else "")
            + (f"; bad glyphs: {bad[:3]}" if bad else ""))


def _html_facts(html_path: Path) -> tuple[list[dict], list[dict], list[list[str]]]:
    from bs4 import BeautifulSoup
    soup = BeautifulSoup(html_path.read_text(encoding="utf-8"), "html.parser")
    top5 = [{k: v for k, v in tr.attrs.items() if k.startswith("data-")} for tr in soup.select("table.top5 tr[data-rank]")]
    cards = [{k: v for k, v in d.attrs.items() if k.startswith("data-")} for d in soup.select("div.card[data-rank]")]
    nums = [sorted(re.findall(r"\d[\d,]*(?:\.\d+)?", d.get_text(" "))) for d in soup.select("div.card[data-rank]")]
    return top5, cards, nums


def parity_checks(paths: dict, rep: Report) -> None:
    """English is the source of truth: the Hindi files must carry exactly the same facts."""
    import pymupdf
    from openpyxl import load_workbook
    from ..reporting.i18n import language, t
    (en_x, en_p, en_t), (hi_x, hi_p, hi_t) = paths["en"], paths["hi"]
    # ---- report: every Top-5 row and Top-10 card, attribute by attribute, plus the numbers printed in each card
    e5, ec, en_nums = _html_facts(en_p.with_suffix(".html"))
    h5, hc, hi_nums = _html_facts(hi_p.with_suffix(".html"))
    diffs = [f"top5 #{a.get('data-rank')}" for a, b in zip(e5, h5) if a != b] + \
            [f"card #{a.get('data-rank')}" for a, b in zip(ec, hc) if a != b]
    ok = len(e5) == len(h5) and len(ec) == len(hc) and not diffs
    rep.add("PASS" if ok else "FAIL", "15. EN = HI: Top 5 and Top 10 (rank, id, URL, rent, maintenance, normalised USD, "
            "total, bedrooms, area, furnished, noise category, risk, confidence, status, score, budget class)",
            f"{len(e5)} + {len(ec)} compared" + (f"; differ: {diffs[:5]}" if diffs else ""))
    num_diff = [i + 1 for i, (a, b) in enumerate(zip(en_nums, hi_nums)) if a != b]
    rep.add("PASS" if not num_diff else "FAIL", "15b. EN = HI: every number printed in each Top-10 card",
            "identical" if not num_diff else f"cards {num_diff}")
    def links(p):
        raw = [ln.get("uri") for page in pymupdf.open(p) for ln in page.get_links() if ln.get("uri")]
        return [u for i, u in enumerate(raw) if i == 0 or u != raw[i - 1]]   # a wrapped link = several rects
    el, hl = links(en_p), links(hi_p)
    rep.add("PASS" if el == hl and el else "FAIL", "15c. EN = HI: PDF hyperlinks identical and in the same order",
            f"{len(el)} links")
    # ---- workbook: every cell of the client sheets; numbers / links / codes equal, text = its Hindi translation
    ew, hw = load_workbook(en_x), load_workbook(hi_x)
    names = SHEET_NAMES["HI"]
    bad = []
    with language("hi"):
        for sheet in ("CLIENT_TOP_PICKS", "EXECUTIVE_SHORTLIST", "ALL_MATCHES", "STRETCH_NEGOTIABLE", "NEAR_MISSES"):
            a, b = ew[sheet], hw[names[sheet]]
            if (a.max_row, a.max_column) != (b.max_row, b.max_column):
                bad.append(f"{sheet}: shape differs")
                continue
            for ra, rb in zip(a.iter_rows(), b.iter_rows()):
                for ca, cb in zip(ra, rb):
                    if (ca.hyperlink.target if ca.hyperlink else None) != (cb.hyperlink.target if cb.hyperlink else None):
                        bad.append(f"{sheet}!{ca.coordinate} link")
                    elif not isinstance(ca.value, str) and ca.value != cb.value:
                        bad.append(f"{sheet}!{ca.coordinate} value")
                    elif isinstance(ca.value, str) and cb.value not in (ca.value, t(ca.value, record=False)) \
                            and not str(cb.value).startswith(ca.value.split(" · ")[0]):
                        bad.append(f"{sheet}!{ca.coordinate} text")
    rep.add("PASS" if not bad else "FAIL", "16. EN = HI workbook: same rows, order, numbers, links and codes; text only "
            "translated", f"{len(bad)} differences" + (f": {bad[:6]}" if bad else ""))
    # ---- contact templates: the same listings and links in the same order
    url_rx = re.compile(r"https?://\S+")
    eu, hu = url_rx.findall(en_t.read_text(encoding="utf-8")), url_rx.findall(hi_t.read_text(encoding="utf-8"))
    rep.add("PASS" if eu == hu else "FAIL", "17. EN = HI contact templates: same links in the same order", f"{len(eu)} links")


def final_qa(ranked: pd.DataFrame, paths: dict, cfg: dict, token: str | None, http,
             production: bool = True) -> tuple[list[str], bool]:
    from ..reporting.i18n import MISSING
    rep = Report()
    urls = []
    for lang in ("en", "hi"):
        tag = lang.upper()
        xlsx, pdf, txt = paths[lang]
        qa_dir = (K.PROCESSED / "qa" if production else Path(pdf).parent / "qa") / lang
        qa_dir.mkdir(parents=True, exist_ok=True)
        u = xlsx_checks(Path(xlsx), rep, production, tag)
        urls = urls or u
        try:
            render_sheet_png(Path(xlsx), SHEET_NAMES[tag]["CLIENT_TOP_PICKS"], qa_dir / "xlsx_client_top_picks.png")
            rep.add("PASS", f"[{tag}] 1b. XLSX rendered for visual review", str(qa_dir / "xlsx_client_top_picks.png"))
        except Exception as exc:  # noqa: BLE001
            rep.add("WARN", f"[{tag}] 1b. XLSX render", f"{type(exc).__name__}: {exc}")
        pdf_links = pdf_checks(Path(pdf), rep, qa_dir, production, tag)
        if production:
            txt_hits = DEMO_RX.findall(Path(txt).read_text(encoding="utf-8"))
            rep.add("PASS" if not txt_hits else "FAIL", f"[{tag}] 6d. No demo data in contact templates",
                    ", ".join(set(txt_hits)))
            rep.add("PASS" if pdf_links else "FAIL", f"[{tag}] 3b. PDF contains clickable links", f"{len(pdf_links)} links")
    devanagari_checks(Path(paths["hi"][1]), rep)
    rep.add("PASS" if not MISSING else "FAIL", "[HI] 18. No untranslated client-facing text",
            f"{len(MISSING)} strings" + (f": {sorted(MISSING)[:6]}" if MISSING else ""))
    parity_checks(paths, rep)
    if production:
        link_checks(urls, http, rep, "3. Top-10 hyperlinks (workbook) respond")
    leak_checks([Path(p) for v in paths.values() for p in v], token, rep)
    rule_checks(ranked, cfg, rep, production)
    if production:
        top_urls = list(ranked[ranked["category"] == "PRIMARY"].sort_values("rank_in_category")
                        .head(cfg["shortlist"]["top_n"])["source_url"])
        link_checks(top_urls, http, rep, "13. Top-10 source URLs re-checked at the end")
    return rep.lines, not rep.failed
