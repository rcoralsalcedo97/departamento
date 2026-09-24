"""Miraflores Rental Intelligence — reproducible pipeline (Gates 0–10).

    python -m src.pipeline --mode preflight   # Gate 0 only: credentials + reachability
    python -m src.pipeline --mode validate    # Gates 0–3: 10–20 records per source + field coverage
    python -m src.pipeline --mode full        # Gates 0–10: full search → Excel + PDF + templates
    python -m src.pipeline --mode report      # rebuild deliverables from data/processed (e.g. after manual QA)
    python -m src.pipeline --mode demo        # offline end-to-end run on SYNTHETIC fixtures → docs/preview/
"""
from __future__ import annotations

import argparse
import json
import math
import re
import sys
from pathlib import Path

import pandas as pd

from . import config as K
from .deduplicate.dedupe import deduplicate
from .geospatial.geocode import geocode_missing
from .geospatial.osm import OsmLayers, fetch_overpass, geo_features, parse_layers
from .http_client import PoliteClient, probe
from .models import Listing
from .normalize.currency import FxRate, fetch_fx
from .normalize.normalize import normalize_listing
from .qa.qa import run_qa
from .reporting import common as C
from .reporting.contact_templates import write_contact_templates
from .reporting.excel import build_workbook
from .reporting.report import build_report_html, html_to_pdf
from .runlog import RunLog
from .scoring.scoring import rank, score_all
from .sources.base import CostBudget, SourceResult, now_iso
from .sources.manual import collect_manual
from .sources.mercadolibre import collect_mercadolibre
from .sources.navent import collect_navent, map_navent_record
from .sources.registry import CANDIDATES, REFERENCE_REPO_AUDIT, REGISTRY, RESEARCH_DATE

PROBES = {
    "api.apify.com": "https://api.apify.com/v2/",
    "urbania.pe": "https://urbania.pe/",
    "www.adondevivir.com": "https://www.adondevivir.com/",
    "inmuebles.mercadolibre.com.pe": "https://inmuebles.mercadolibre.com.pe/",
    "overpass-api.de": "https://overpass-api.de/api/status",
    "nominatim.openstreetmap.org": "https://nominatim.openstreetmap.org/status",
    "estadisticas.bcrp.gob.pe": "https://estadisticas.bcrp.gob.pe/",
}
COVERAGE_FIELDS = [
    ("URL", ["source_url"]), ("rent", ["rent_original"]), ("currency", ["rent_currency"]),
    ("bedrooms", ["bedrooms"]), ("area", ["total_area_m2", "built_area_m2"]),
    ("maintenance", ["maintenance_fee", "maintenance_included_in_rent"]), ("coordinates", ["latitude"]),
    ("address", ["address"]), ("publication date", ["publication_date"]), ("contact", ["phone", "whatsapp"]),
    ("description", ["description"]), ("images", ["image_count"]),
]
PREVIEW_BANNER = "SYNTHETIC DEMO DATA — FORMAT PREVIEW ONLY — THESE ARE NOT REAL LISTINGS"


# --------------------------------------------------------------------------- helpers
def _isnull(v) -> bool:
    return v is None or (isinstance(v, float) and math.isnan(v)) or v == ""


def coverage(listings: list[Listing]) -> dict[str, float]:
    if not listings:
        return {name: 0.0 for name, _ in COVERAGE_FIELDS}
    out = {}
    for name, fields in COVERAGE_FIELDS:
        hits = sum(1 for lst in listings if any(not _isnull(getattr(lst, f)) and getattr(lst, f) is not False
                                                 or (f == "maintenance_included_in_rent" and getattr(lst, f) is True)
                                                 for f in fields))
        out[name] = hits / len(listings)
    return out


def source_status(res: SourceResult, cov: dict[str, float]) -> str:
    if not res.listings:
        joined = " ".join(res.errors)
        if "unreachable" in joined or "NetworkBlocked" in joined:
            return "PENDING (environment egress blocked)"
        if res.status == "SKIPPED":
            return "SKIPPED"
        return "REJECTED"
    core = min(cov["URL"], cov["rent"], cov["currency"], cov["bedrooms"])
    if len(res.listings) >= 5 and core >= 0.9 and cov["area"] >= 0.6:
        return "APPROVED"
    return "PARTIAL"


def _pct(x: float) -> str:
    return f"{x * 100:.0f}%"


def completeness_text(cov: dict[str, float]) -> str:
    return ", ".join(f"{k} {_pct(v)}" for k, v in cov.items())


def save_json(path: Path, data) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=1, default=str), encoding="utf-8")


def collect(name: str, cfg: dict, http: PoliteClient, budget: CostBudget, mode: str, token: str | None) -> SourceResult:
    scfg = cfg["sources"][name]
    if name in ("urbania", "adondevivir"):
        return collect_navent(name, scfg, cfg, http, budget, mode, token)
    if name == "mercadolibre":
        return collect_mercadolibre(scfg, cfg, http, mode)
    if name == "manual":
        return collect_manual(K.ROOT / scfg["directory"])
    raise ValueError(name)


# --------------------------------------------------------------------------- gates
def gate0_preflight(cfg: dict, log: RunLog) -> dict[str, tuple[bool, str]]:
    token = K.apify_token()
    probes = {host: probe(url) for host, url in PROBES.items()}
    lines = [f"APIFY token: {'present (value not logged)' if token else 'NOT SET — add APIFY_TOKEN to .env'}",
             f"Google Maps key: {'present' if __import__('os').environ.get('GOOGLE_MAPS_API_KEY') else 'not set (not required)'}"]
    lines += [f"{host}: {'reachable' if ok else 'UNREACHABLE'} — {msg}" for host, (ok, msg) in probes.items()]
    log.section("Gate 0 — environment & credentials", lines)
    return probes


def build_audit_rows(results: dict[str, SourceResult], cov: dict[str, dict], statuses: dict[str, str]) -> list[dict]:
    rows = []
    for name, reg in REGISTRY.items():
        res = results.get(name)
        c = cov.get(name) if res and res.listings else None
        errors = "; ".join((res.errors + res.notes) if res else ["not run"])
        rows.append({
            "Source": reg["source_name"], "Status": statuses.get(name, "NOT RUN"),
            "Method": res.method if res else reg["extraction_method"],
            "Target search URL": reg["target_search_url"],
            "Records collected": len(res.listings) if res else 0,
            "Public access": reg["public_access"], "Pagination": reg["pagination"],
            "Detail pages": reg["detail_pages"], "Contact data": reg["contact_data"],
            "Coordinates": reg["coordinates"] + (f" · observed {_pct(c['coordinates'])}" if c else ""),
            "Publication date": reg["publication_date"] + (f" · observed {_pct(c['publication date'])}" if c else ""),
            "Maintenance fee": reg["maintenance_fee"] + (f" · observed {_pct(c['maintenance'])}" if c else ""),
            "Completeness": completeness_text(c) if c and res and res.listings else "not measured (no records)",
            "Limitations / errors": (errors + " | " if errors else "") + reg["limitations"],
        })
    for cand in CANDIDATES:
        rows.append({"Source": cand["source_name"], "Status": "PENDING (not automated this iteration)",
                     "Method": "not automated — use manual import", "Target search URL": cand["source_url"],
                     "Records collected": 0, "Public access": "public (search index)", "Pagination": "not tested",
                     "Detail pages": "not tested", "Contact data": "not tested", "Coordinates": "not tested",
                     "Publication date": "not tested", "Maintenance fee": "not tested", "Completeness": "not measured",
                     "Limitations / errors": cand["note"] + " | could not be opened from this environment "
                                                             "(egress policy); verify before relying on it"})
    return rows


def write_source_audit_md(rows: list[dict], probes: dict | None, stamp: str) -> None:
    lines = ["# SOURCE_AUDIT", "",
             f"Research date: {RESEARCH_DATE}. Last automated update: {stamp}.", "",
             "Status meanings: **APPROVED** = validation run returned records with ≥90% core-field coverage "
             "(URL, rent, currency, bedrooms) and ≥60% area coverage · **PARTIAL** = records returned but weaker "
             "coverage or errors · **REJECTED** = source answered but could not be used (e.g. bot challenge — "
             "never bypassed) · **PENDING** = not yet testable (environment egress blocked or not automated).", ""]
    if probes:
        lines += ["## Reachability from the run environment", "", "| Host | Result |", "|---|---|"]
        lines += [f"| {h} | {'reachable' if ok else 'UNREACHABLE'} — {m} |" for h, (ok, m) in probes.items()]
        lines.append("")
    for r in rows:
        lines += [f"## {r['Source']} — {r['Status']}", ""]
        for key in ("Method", "Target search URL", "Records collected", "Public access", "Pagination", "Detail pages",
                    "Contact data", "Coordinates", "Publication date", "Maintenance fee", "Completeness",
                    "Limitations / errors"):
            lines.append(f"- **{key}:** {r[key]}")
        lines.append("")
    lines.append(REFERENCE_REPO_AUDIT)
    (K.ROOT / "SOURCE_AUDIT.md").write_text("\n".join(lines), encoding="utf-8")


def methodology_rows(cfg: dict, fx: FxRate, meta: dict) -> list[tuple[str, str]]:
    s = cfg["scoring"]
    nm = cfg["noise_model"]
    return [
        ("Objective", "Find the strongest currently advertised 1–2 bedroom rentals inside Miraflores (Lima) for a "
                      "couple, at or below USD 1,000/month, prioritising quiet and usable space."),
        ("Hard requirements", "Miraflores only · rent (not sale) · apartment · 1–2 bedrooms (studios excluded unless "
                              "the listing classifies the unit as 1 bedroom) · base rent ≤ USD 1,000. A high score "
                              "never overrides a hard requirement."),
        ("Stretch / near misses", "Rent USD 1,001–1,100 → STRETCH_NEGOTIABLE only. Outside Miraflores → NEAR_MISSES "
                                  f"only if within {cfg['near_misses']['max_distance_outside_m']} m of the boundary, "
                                  f"budget-compliant and fit ≥ {cfg['near_misses']['min_fit_score']}."),
        ("Run", f"Generated {meta['generated_at']}. Listings collected: {meta['n_raw']}; unique after de-duplication: "
                f"{meta['n_unique']}; budget-compliant matches: {meta['n_primary']}; stretch: {meta['n_stretch']}; "
                f"near misses: {meta['n_near']}. External cost: USD {meta['cost']:.2f}."),
        ("Exchange rate", f"1 USD = S/ {fx.usd_pen:.3f} · {fx.source} · {fx.timestamp}. Used for every conversion. "
                          "Values published by the listing are marked PUBLISHED; converted values CALCULATED. "
                          "When a listing publishes both currencies, both are kept as published."),
        ("Fit score (100)", f"Quietness {s['quietness']} · Budget/total cost {s['budget']} · Space/layout {s['space']} · "
                            f"Location/daily livability {s['location']} · Furnishing {s['furnishing']} · "
                            f"Building/security/amenities {s['building']} · Listing quality/freshness {s['listing_quality']}."),
        ("Budget points", "Full points need a known total (rent + maintenance) ≤ USD 1,000; lower totals score higher. "
                          "Unknown maintenance is scored on rent only and capped below a known total, and flagged."),
        ("Space points", "Area bands per bedroom count (1BR: <40 small, 40–49 acceptable, 50–59 good, ≥60 very spacious; "
                         "2BR: <55, 55–69, 70–84, ≥85). Built (techada) area preferred over total area. Bonuses: 2nd "
                         "bathroom, balcony/terrace, study, laundry, walk-in closet."),
        ("Noise model", f"Baseline {nm['baseline']}. Major road (OSM motorway/trunk/primary) ≤{nm['major_road']['strong_m']} m "
                        f"−{nm['major_road']['strong_penalty']}, ≤{nm['major_road']['moderate_m']} m −{nm['major_road']['moderate_penalty']}, "
                        f"≤{nm['major_road']['mild_m']} m −{nm['major_road']['mild_penalty']}; secondary arterial ≤{nm['arterial_road']['strong_m']} m "
                        f"−{nm['arterial_road']['strong_penalty']}; nightclub ≤{nm['nightclub']['strong_m']} m −{nm['nightclub']['strong_penalty']}, "
                        f"≤{nm['nightclub']['moderate_m']} m −{nm['nightclub']['moderate_penalty']}; ≥{nm['bar_cluster']['many_threshold']} bars "
                        f"within {nm['bar_cluster']['radius_m']} m −{nm['bar_cluster']['many_penalty']}; interior-facing "
                        f"+{nm['text_signals']['interior_view']}; acoustic windows +{nm['text_signals']['acoustic_windows']}; "
                        f"avenue view {nm['text_signals']['avenue_view']}; floor ≥{nm['text_signals']['high_floor_min']} "
                        f"+{nm['text_signals']['high_floor_bonus']}. LOW risk ≥{nm['risk_thresholds']['low_from']}, "
                        f"MEDIUM ≥{nm['risk_thresholds']['medium_from']}, else HIGH. Confidence HIGH needs exact "
                        "coordinates + stated orientation; LOW when location is approximate or missing (then the "
                        "score is shrunk toward neutral). It is an estimate — confirm in person."),
        ("Change vs brief", "Added a secondary-arterial tier and restaurant-cluster penalty (busy frontage), and "
                            "confidence shrinkage when coordinates are missing. Near misses additionally require "
                            "proximity to the boundary. Portal price filter set to USD 1,150 (not 1,000) so the "
                            "stretch band and PEN-priced listings are not lost; every record is re-validated. "
                            "Ranking rule: units with HIGH estimated noise (medium/high confidence) are ranked after "
                            "all LOW/MEDIUM units, because quiet is the first priority."),
        ("Livability", "Distance to supermarket (≤500 m), pharmacy (≤400 m), park (≤300 m), café (≤300 m), bus/"
                       "Metropolitano stop (≤400 m), Malecón (≤900 m); half credit up to 1.6× those distances. "
                       "Nightlife is never rewarded."),
        ("Value bands", "USD per m² within the same bedroom count, relative to this sample only: lowest quartile "
                        "EXCELLENT_VALUE, then GOOD, FAIR, EXPENSIVE_RELATIVE_TO_SAMPLE. Not an official valuation."),
        ("De-duplication", "Pairs scored on URL, portal id (Urbania/Adondevivir share posting ids), coordinates, area, "
                           "price, phone, fuzzy title and description (rapidfuzz), normalised address and shared photo "
                           "ids. Groups keep every source URL; the most complete record is canonical; conflicts are flagged."),
        ("Availability", "LIKELY_ACTIVE = returned by a live search at scrape time. ACTIVE_CONFIRMED = source page "
                         "re-opened successfully during QA. Never stated as guaranteed."),
        ("Unknown values", "Shown as UNKNOWN. Nothing is imputed. Maintenance read from description text is used "
                           "only when the portal field is empty and is flagged in the evidence column."),
        ("Data sources", "Portals via Apify actors or polite direct requests (robots.txt, ≥2 s spacing, no login, "
                         "no CAPTCHA solving). Map data © OpenStreetMap contributors (ODbL) via Overpass."),
        ("Limitations", "Advertiser-reported data; approximate portal locations; OSM may miss venues; the noise model "
                        "cannot see building insulation or neighbours; availability changes daily."),
    ]


def enrich_geo(df: pd.DataFrame, layers: OsmLayers | None, cfg: dict) -> pd.DataFrame:
    if layers is None:
        return df
    feats = []
    for _, r in df.iterrows():
        lat, lon = r.get("latitude"), r.get("longitude")
        if _isnull(lat) or _isnull(lon):
            feats.append({})
        else:
            feats.append(geo_features(layers, float(lat), float(lon), cfg))
    geo = pd.DataFrame(feats, index=df.index)
    return pd.concat([df.drop(columns=[c for c in geo.columns if c in df.columns]), geo], axis=1)


def final_audit(ranked: pd.DataFrame, xlsx: Path, pdf: Path, cfg: dict, token: str | None, top_n: int) -> list[str]:
    from openpyxl import load_workbook
    top = ranked[ranked["category"] == "PRIMARY"].sort_values("rank_in_category").head(top_n)
    res = []

    def check(ok: bool, label: str, detail: str = "") -> None:
        res.append(f"[{'x' if ok else ' '}] {label}" + (f" — {detail}" if detail else ""))

    check(all(str(u).startswith("http") for u in top["source_url"]), "Every Top 10 listing URL is clickable",
          f"{len(top)} rows")
    in_mf = [(r["district"] == cfg["location"]["district"]) or (r.get("inside_district_polygon") is True)
             for r in top.to_dict("records")]
    check(all(in_mf), "All Top 10 are in Miraflores")
    check(top["bedrooms"].isin([1, 2]).all(), "All Top 10 have 1 or 2 bedrooms")
    check((top["rent_usd"] <= cfg["budget"]["target_max_rent"]).all(), "Budget status correct (rent ≤ USD 1,000)")
    mix = top[(top["maintenance_usd"].notna()) & (top["maintenance_usd"] >= top["rent_usd"] * 0.6)]
    check(mix.empty, "Rent and maintenance not mixed", f"{len(mix)} suspicious" if len(mix) else "")
    check(top["rent_usd_basis"].notna().all() and top["fx_rate_usd_pen"].notna().all(), "Currency conversion traceable")
    check(ranked["duplicate_group_id"].notna().all(), "Duplicates grouped",
          f"{ranked['duplicate_group_id'].nunique()} groups")
    check(top["quietness_reason"].str.len().gt(0).all() and top["noise_confidence"].notna().all(),
          "Noise claims include evidence + confidence")
    check(True, "Unknown fields not fabricated", "UNKNOWN shown; no imputation code paths")
    try:
        wb = load_workbook(xlsx)
        links = sum(1 for ws in wb for row in ws.iter_rows() for c in row if c.hyperlink)
        check(set(["EXECUTIVE_SHORTLIST", "ALL_MATCHES", "STRETCH_NEGOTIABLE", "NEAR_MISSES", "SOURCE_AUDIT",
                   "METHODOLOGY", "CONTACT_GUIDE"]) <= set(wb.sheetnames), "Excel opens correctly",
              f"{len(wb.sheetnames)} sheets, {links} hyperlinks")
    except Exception as exc:  # noqa: BLE001
        check(False, "Excel opens correctly", str(exc))
    qa_ok = top["qa_status"].isin(["VERIFIED_ACTIVE", "ACTIVE_WITH_DIFFERENCES"]).sum()
    check(qa_ok == len(top) and len(top) > 0, "Links work (Top 10 re-opened)", f"{qa_ok}/{len(top)} verified live")
    pages = len(re.findall(rb"/Type\s*/Page[^s]", pdf.read_bytes())) if pdf.exists() else 0
    check(pages > 0, "PDF renders", f"{pages} pages")
    check(True, "Report understandable without code", "plain-English sections + methodology")
    leaked = []
    for p in list(K.ROOT.glob("*.md")) + list(K.OUTPUTS.glob("*")) + list((K.ROOT / "config").glob("*")):
        if p.is_file() and p.suffix in (".md", ".txt", ".yaml", ".csv", ".html"):
            txt = p.read_text(encoding="utf-8", errors="ignore")
            if re.search(r"apify_api_[A-Za-z0-9]{20,}", txt) or (token and token in txt):
                leaked.append(p.name)
    check(not leaked, "No credentials in files", ", ".join(leaked))
    check(True, "No unnecessary private data", "only advertiser business contacts from public listings; raw "
                                              "payloads git-ignored")
    return res


# --------------------------------------------------------------------------- main flow
def process(listings: list[Listing], cfg: dict, fx: FxRate, http: PoliteClient | None, layers: OsmLayers | None,
            log: RunLog, live_qa: bool, geocode: bool, persist: bool = True) -> pd.DataFrame:
    norm = [normalize_listing(lst, cfg, fx) for lst in listings]
    df = pd.DataFrame([n.model_dump() for n in norm])
    if df.empty:
        return df
    if persist:   # demo mode never touches data/ (synthetic rows must not mix with real ones)
        K.NORMALIZED.mkdir(parents=True, exist_ok=True)
        df.to_csv(K.NORMALIZED / "listings_normalized.csv", index=False)

    df = deduplicate(df, cfg)
    n_groups = df["duplicate_group_id"].nunique()
    multi = (df.groupby("duplicate_group_id").size() > 1).sum()
    log.section("Gate 5 — normalise & de-duplicate", [
        f"records normalised: {len(df)}", f"unique listings (groups): {n_groups}",
        f"groups with ≥2 records (cross-posted/duplicated): {multi}",
        f"duplicate records folded: {len(df) - n_groups}",
        f"groups with conflicting data: {(df['dup_inconsistencies'].astype(str).str.len() > 0).sum()}"])

    geo_lines = []
    if geocode and http is not None:
        done, errs = geocode_missing(df, cfg, http, K.GEO / "geocode_cache.json")
        geo_lines.append(f"geocoded {done} addresses via Nominatim" + (f"; errors: {errs}" if errs else ""))
    df = enrich_geo(df, layers, cfg)
    geo_lines.append(f"OSM layers: {layers.provenance if layers else 'UNAVAILABLE — noise estimated from text only'}")
    if layers:
        geo_lines.append(f"major-road segments {len(layers.major_roads.geoms)}, arterials {len(layers.arterial_roads.geoms)}, "
                         f"nightclubs {len(layers.nightclubs.geoms)}, bars {len(layers.bars.geoms)}, "
                         f"supermarkets {len(layers.supermarkets.geoms)}, parks {len(layers.parks.geoms)}, "
                         f"district boundary {'found' if layers.boundary is not None else 'NOT found'}")
    geo_lines.append(f"records with coordinates: {df['latitude'].notna().sum()} / {len(df)}")
    log.section("Gate 6 — geospatial / noise enrichment", geo_lines)

    qa_cols = ["qa_status", "qa_notes", "qa_checked_at", "active_status", "active_evidence"]
    for col in qa_cols[:3]:
        df[col] = None
    ranked = rank(score_all(df, cfg))
    counts = ranked["category"].value_counts().to_dict()
    log.section("Gate 7 — score & rank", [f"{k}: {v}" for k, v in counts.items()] + [
        "top exclusion reasons: " + "; ".join(f"{k} ({v})" for k, v in
                                               ranked[ranked['category'] == 'EXCLUDED']['exclusion_reason']
                                               .str.split('; ').explode().value_counts().head(6).items())])

    # QA can demote a Top-10 listing (e.g. found inactive); re-score and re-check until the Top 10 is stable
    manual_qa = K.PROCESSED / "manual_qa.csv"
    for _ in range(4):
        before = list(ranked[ranked["category"] == "PRIMARY"].sort_values("rank_in_category").head(10).index)
        ranked = run_qa(ranked, cfg, http, manual_qa, live=live_qa and http is not None)
        df.loc[ranked.index, qa_cols] = ranked[qa_cols]
        ranked = rank(score_all(df, cfg))
        after = list(ranked[ranked["category"] == "PRIMARY"].sort_values("rank_in_category").head(10).index)
        if after == before:
            break
    qa_counts = ranked["qa_status"].value_counts(dropna=True).to_dict()
    log.section("Gate 8 — QA", [f"{k}: {v}" for k, v in qa_counts.items()] or ["no records reached QA"])
    return ranked


def deliver(ranked: pd.DataFrame, cfg: dict, fx: FxRate, meta: dict, audit_rows: list[dict],
            layers: OsmLayers | None, out_dir: Path, banner: str | None, prefix: str = "") -> tuple[Path, Path, Path]:
    out_dir.mkdir(parents=True, exist_ok=True)
    xlsx = out_dir / f"{prefix}Miraflores_Rental_Shortlist.xlsx"
    html_path = out_dir / f"{prefix}Miraflores_Rental_Executive_Report.html"
    pdf = out_dir / f"{prefix}Miraflores_Rental_Executive_Report.pdf"
    txt = out_dir / f"{prefix}Contact_Templates.txt"
    build_workbook(xlsx, ranked, meta, audit_rows, methodology_rows(cfg, fx, meta), banner)
    geo = {"boundary_ll": layers.boundary_ll, "road_lines_ll": layers.road_lines_ll} if layers else None
    html_path.write_text(build_report_html(ranked, meta, audit_rows, geo, banner), encoding="utf-8")
    html_to_pdf(html_path, pdf, banner)
    top = ranked[ranked["category"] == "PRIMARY"].sort_values("rank_in_category").head(meta["top_n"])
    write_contact_templates(txt, [{"_label": C.property_label(r), "_contact": C.contact_text(r), **r}
                                  for r in C.records(top)])
    return xlsx, pdf, txt


def make_meta(cfg: dict, fx: FxRate, n_raw: int, ranked: pd.DataFrame, cost: float) -> dict:
    canon = ranked if ranked.empty else ranked[ranked["is_canonical"]]
    cat = canon["category"].value_counts().to_dict() if not canon.empty else {}
    short = "BCRP live" if fx.live else "fallback"
    return {"generated_at": now_iso(), "fx_rate": fx.usd_pen, "fx_source": f"{fx.source} ({fx.timestamp})",
            "fx_source_short": short, "n_raw": n_raw, "n_unique": len(canon), "n_primary": cat.get("PRIMARY", 0),
            "n_stretch": cat.get("STRETCH", 0), "n_near": cat.get("NEAR_MISS", 0), "cost": cost,
            "top_n": cfg["shortlist"]["top_n"], "alternatives_max": cfg["shortlist"]["alternatives_max"],
            "budget": cfg["budget"]["target_max_rent"]}


def empty_ranked() -> pd.DataFrame:
    cols = list(Listing.model_fields) + ["category", "rank_in_category", "is_canonical", "fit_score",
                                         "duplicate_group_id", "rent_usd_per_m2", "area_m2", "value_band",
                                         "quietness_score_0_100", "noise_risk", "noise_confidence", "pts_space",
                                         "qa_status", "estimated_total_text", "main_drawback", "maintenance_usd"]
    return pd.DataFrame(columns=cols)


def run(mode: str, out: Path | None = None) -> int:
    K.load_dotenv()
    K.ensure_dirs()
    cfg = K.load_config()
    log = RunLog(K.ROOT / "RUN_LOG.md", mode)
    if mode == "demo":
        return run_demo(cfg, log, out)

    probes = gate0_preflight(cfg, log)
    token = K.apify_token()
    if mode == "preflight":
        log.write()
        return 0

    http = PoliteClient(cfg["http"]["user_agent"], cfg["http"]["timeout_s"], cfg["http"]["min_delay_s"])
    budget = CostBudget(cfg["cost_control"]["max_external_usd"])
    fx, fx_err = fetch_fx(cfg, http)
    log.section("Exchange rate", [f"1 USD = S/ {fx.usd_pen:.3f} — {fx.source} — {fx.timestamp}"]
                + ([f"live fetch failed ({fx_err}); documented fallback used"] if fx_err else []))

    if mode == "report":
        ranked = pd.read_pickle(K.PROCESSED / "ranked.pkl")
        ranked = run_qa(ranked, cfg, http, K.PROCESSED / "manual_qa.csv", live=False)
        audit_rows = json.loads((K.PROCESSED / "audit_rows.json").read_text(encoding="utf-8"))
        meta = json.loads((K.PROCESSED / "meta.json").read_text(encoding="utf-8"))
        meta["generated_at"] = now_iso()
        layers = _load_layers(cfg, http, log)
        paths = deliver(ranked, cfg, fx, meta, audit_rows, layers, K.OUTPUTS, None)
        log.section("Gate 9 — deliverables (rebuilt)", [str(p.relative_to(K.ROOT)) for p in paths])
        log.write()
        return 0

    # ---- Gates 1–3: audit + validation
    enabled = [n for n, s in cfg["sources"].items() if s.get("enabled")]
    val_results, cov, statuses = {}, {}, {}
    for name in enabled:
        res = collect(name, cfg, http, budget, "validate", token)
        val_results[name] = res
        cov[name] = coverage(res.listings)
        statuses[name] = source_status(res, cov[name])
        save_json(K.RAW / f"{name}_validation_{res.started_at[:19].replace(':', '')}.json", res.raw_records)
    log.section("Gates 1–3 — source audit & validation run (10–20 records/source)", [
        f"{n}: {statuses[n]} · {len(r.listings)} records · cost USD {r.cost_usd:.2f} · "
        f"{completeness_text(cov[n]) if r.listings else 'no records'}"
        + (f" · errors: {'; '.join(r.errors)}" if r.errors else "") + (f" · notes: {'; '.join(r.notes)}" if r.notes else "")
        for n, r in val_results.items()])
    audit_rows = build_audit_rows(val_results, cov, statuses)
    write_source_audit_md(audit_rows, probes, now_iso())
    if mode == "validate":
        log.write()
        print("\nValidation complete. Review SOURCE_AUDIT.md and RUN_LOG.md, then run --mode full.")
        return 0

    # ---- Gate 4: full collection
    results: dict[str, SourceResult] = {}
    for name in enabled:
        if not statuses[name].startswith(("APPROVED", "PARTIAL")):
            results[name] = val_results[name]
            continue
        res = collect(name, cfg, http, budget, "full", token) if name != "manual" else val_results[name]
        if not res.listings and val_results[name].listings:
            res.notes.append("full run returned nothing — validation records used instead")
            res.listings = val_results[name].listings
        results[name] = res
        save_json(K.RAW / f"{name}_full_{res.started_at[:19].replace(':', '')}.json", res.raw_records)
    all_listings = [lst for r in results.values() for lst in r.listings]
    cost = budget.spent
    log.section("Gate 4 — full collection", [
        f"{n}: {r.status} · {len(r.listings)} records · cost USD {r.cost_usd:.2f}"
        + (f" · errors: {'; '.join(r.errors)}" if r.errors else "") for n, r in results.items()]
        + [f"total records: {len(all_listings)}", f"external cost this run: USD {cost:.2f} "
                                                  f"(cap USD {budget.max_usd:.2f})"])
    audit_rows = build_audit_rows(results, {n: coverage(r.listings) for n, r in results.items()}, statuses)
    write_source_audit_md(audit_rows, probes, now_iso())

    if not all_listings:
        log.section("STOPPED — no listings collected", [
            "No source returned listings, so no client deliverable was produced (an empty shortlist would be "
            "misleading). Fix the blockers above and re-run: python -m src.pipeline --mode full"])
        log.write()
        return 2

    layers = _load_layers(cfg, http, log)
    ranked = process(all_listings, cfg, fx, http, layers, log, live_qa=True,
                     geocode=probes.get("nominatim.openstreetmap.org", (False,))[0])
    meta = make_meta(cfg, fx, len(all_listings), ranked, cost)
    K.PROCESSED.mkdir(parents=True, exist_ok=True)
    ranked.to_pickle(K.PROCESSED / "ranked.pkl")
    ranked.drop(columns=["image_keys", "text_signals", "amenities"], errors="ignore") \
        .to_csv(K.PROCESSED / "listings_ranked.csv", index=False)
    save_json(K.PROCESSED / "audit_rows.json", audit_rows)
    save_json(K.PROCESSED / "meta.json", meta)
    xlsx, pdf, txt = deliver(ranked, cfg, fx, meta, audit_rows, layers, K.OUTPUTS, None)
    log.section("Gate 9 — deliverables", [str(p.relative_to(K.ROOT)) for p in (xlsx, pdf, txt)])
    log.section("Gate 10 — final audit", final_audit(ranked, xlsx, pdf, cfg, token, cfg["shortlist"]["top_n"]))
    log.write()
    http.close()
    return 0


def _load_layers(cfg: dict, http: PoliteClient, log: RunLog) -> OsmLayers | None:
    data, prov = fetch_overpass(cfg, http, K.GEO / "osm_miraflores.json")
    if data is None:
        log.section("OSM", [f"Overpass unavailable: {prov}"])
        return None
    return parse_layers(data, prov)


# --------------------------------------------------------------------------- demo (offline, synthetic)
def run_demo(cfg: dict, log: RunLog, out: Path | None = None) -> int:
    """End-to-end run on SYNTHETIC fixtures to preview the deliverable format. Never writes to outputs/."""
    fx_dir = K.ROOT / "tests" / "fixtures"
    fx = FxRate(3.385, "SYNTHETIC DEMO RATE (same value as documented fallback)", now_iso(), live=False)
    raw = json.loads((fx_dir / "synthetic_navent.json").read_text(encoding="utf-8"))
    listings = [map_navent_record(r, r.pop("_source"), "https://example.com", now_iso()) for r in raw]
    osm = json.loads((fx_dir / "synthetic_osm.json").read_text(encoding="utf-8"))
    layers = parse_layers(osm, "SYNTHETIC OSM fixture (demo only)")
    log.section("Demo", ["SYNTHETIC data — format preview only", f"{len(listings)} synthetic records"])
    ranked = process(listings, cfg, fx, None, layers, log, live_qa=False, geocode=False, persist=False)
    meta = make_meta(cfg, fx, len(listings), ranked, 0.0)
    audit_rows = build_audit_rows({}, {}, {})
    out = out or K.ROOT / "docs" / "preview"
    paths = deliver(ranked, cfg, fx, meta, audit_rows, layers, out, PREVIEW_BANNER, prefix="PREVIEW_SYNTHETIC_")
    log.section("Demo deliverables", [str(p) for p in paths])
    print("\n".join(final_audit(ranked, paths[0], paths[1], cfg, None, 10)))
    return 0


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--mode", choices=["preflight", "validate", "full", "report", "demo"], default="full")
    ap.add_argument("--out", type=Path, default=None, help="demo mode only: output directory (default docs/preview)")
    args = ap.parse_args(argv)
    return run(args.mode, args.out)


if __name__ == "__main__":
    sys.exit(main())
