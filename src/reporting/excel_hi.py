"""Hindi workbook = the finished English workbook, translated cell by cell.

Nothing is recomputed: the file is copied from the English master and only human-readable
text is translated. Numbers, hyperlinks, IDs, prices, addresses, names and classification codes
are left exactly as they are, so the two workbooks cannot disagree on facts.
"""
from __future__ import annotations

import re
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment

from .contact_templates import TEMPLATE_HI
from .i18n import code_hi, codes_hi, language, t

SHEETS = {"CLIENT_TOP_PICKS": "शीर्ष विकल्प", "EXECUTIVE_SHORTLIST": "कार्यकारी शॉर्टलिस्ट",
          "ALL_MATCHES": "सभी उपयुक्त विकल्प", "STRETCH_NEGOTIABLE": "स्ट्रेच (मोलभाव योग्य)",
          "NEAR_MISSES": "लगभग उपयुक्त व सीमा-रेखा", "SOURCE_AUDIT": "स्रोत ऑडिट",
          "METHODOLOGY": "कार्यप्रणाली", "CONTACT_GUIDE": "संपर्क मार्गदर्शिका"}
# values that are provenance or identity: shown exactly as published / as recorded
KEEP_COLS = {"Property", "Source", "Agency / agent", "Phone", "Duplicate group", "Rent as published", "Published",
             "Target search URL", "Duplicate conflicts", "QA notes", "Rank", "# Sources"}
# classification codes, explained in Hindi on the METHODOLOGY sheet
CODE_COLS = {"Budget Class", "Category", "Rent USD basis", "Maintenance basis", "Area basis", "Space band",
             "Value band", "Active status", "Foreign Tenant", "Red flags", "QA Status", "Status", "Location source"}
AUDIT_NOTE = ("इस शीट का तकनीकी स्रोत-विवरण मूल (अंग्रेज़ी) रूप में रखा गया है ताकि स्रोत की प्रामाणिकता बनी रहे; "
              "केवल शीर्षक अनुवादित हैं।")


def _header_row(ws) -> int | None:
    for r in range(1, min(ws.max_row, 12) + 1):
        c = ws.cell(r, 1)
        if c.fill is not None and c.fill.fill_type == "solid" and str(c.fill.fgColor.rgb or "").endswith("1F3A4D"):
            return r
    return None


def _translate_formula(f: str) -> str:
    def left(m):
        hi = t(m.group(3))
        return f'LEFT({m.group(1)},{len(hi)})="{hi}"'
    f = re.sub(r'LEFT\((\$?[A-Z]+\$?\d+),(\d+)\)="([^"]+)"', left, f)
    return re.sub(r'"([^"]+)"', lambda m: f'"{code_hi(m.group(1)) if code_hi(m.group(1)) != m.group(1) else t(m.group(1), record=False)}"', f)


def _wrap_if_long(ws, cell, text: str) -> None:
    width = ws.column_dimensions[cell.column_letter].width or 8.43
    if len(text) > width * 0.95:
        cell.alignment = Alignment(wrap_text=True, vertical=(cell.alignment.vertical if cell.alignment else "top") or "top")


def translate_workbook(en_path: Path, hi_path: Path, methodology_hi: list[tuple[str, str]]) -> Path:
    wb = load_workbook(en_path)
    with language("hi"):
        for ws in wb.worksheets:
            name = ws.title
            hdr = _header_row(ws)
            headers = {c.column: str(c.value) for c in ws[hdr]} if hdr else {}
            for row in ws.iter_rows():
                for c in row:
                    if c.comment is not None and c.comment.text:
                        c.comment = Comment(t(c.comment.text), c.comment.author or "pipeline")
                    v = c.value
                    if not isinstance(v, str):
                        continue
                    if name == "METHODOLOGY" and c.row >= 5:
                        continue                  # replaced below with the Hindi methodology
                    if name == "CONTACT_GUIDE" and c.row in (5, 6):
                        continue                  # Spanish message kept; English meaning replaced below
                    col = headers.get(c.column)
                    if hdr and c.row > hdr:
                        if col in CODE_COLS:
                            c.value = codes_hi(v)     # "हिंदी (CODE)": natural label, code kept for traceability
                            if c.value != v:
                                _wrap_if_long(ws, c, c.value)
                            continue
                        if name == "SOURCE_AUDIT" or col in KEEP_COLS:
                            continue
                        if col in ("Who to contact", "WhatsApp / Contact"):
                            c.value = " · ".join(t(p, record=False) for p in v.split(" · "))
                        elif col == "Location":
                            if " · " not in v:
                                continue          # an address or zone as published
                            head, rest = v.split(" · ", 1)
                            c.value = f"{head} · {t(rest)}"
                        else:
                            c.value = t(v)
                    else:
                        c.value = t(v)            # titles, subtitles, headers
                    if c.value != v and hdr and c.row > hdr:
                        _wrap_if_long(ws, c, c.value)   # titles/subtitles spill across empty cells, as in English
            for rng in ws.conditional_formatting:
                for rule in rng.rules:
                    if rule.formula:
                        rule.formula = [_translate_formula(f) for f in rule.formula]
            if name == "SOURCE_AUDIT":
                ws.cell(2, 1).value = f"{ws.cell(2, 1).value} · {AUDIT_NOTE}"
            if name == "METHODOLOGY":
                for i, (k, v) in enumerate(methodology_hi, start=5):
                    ws.cell(i, 1).value, ws.cell(i, 2).value = k, v
            if name == "CONTACT_GUIDE":
                ws.cell(5, 1).value = "भेजने के लिए संदेश (स्पेनिश)"
                ws.cell(6, 1).value = "संदेश का अर्थ (हिंदी)"
                ws.cell(6, 2).value = TEMPLATE_HI
            ws.title = SHEETS.get(name, name)
    hi_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(hi_path)
    return hi_path
