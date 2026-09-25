"""Client-facing workbook: outputs/Miraflores_Rental_Shortlist.xlsx"""
from __future__ import annotations

from pathlib import Path
from typing import Any, Callable

import pandas as pd
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import common as C
from .contact_templates import TEMPLATE_EN, TEMPLATE_ES

NAVY = "1F3A4D"
HEADER_FONT = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
BODY_FONT = Font(name="Calibri", size=10, color="1B1B1B")
LINK_FONT = Font(name="Calibri", size=10, color="1F5FBF", underline="single")
TITLE_FONT = Font(name="Calibri", size=14, bold=True, color=NAVY)
SUB_FONT = Font(name="Calibri", size=9, italic=True, color="52514E")
WARN_FONT = Font(name="Calibri", size=11, bold=True, color="B42318")
THIN = Side(style="thin", color="D9DCE1")
BORDER = Border(bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
TOP = Alignment(vertical="top")

USD_FMT = '"USD" #,##0'
PPM2_FMT = '0.0'

Col = tuple[str, Callable[[dict], Any], float, str | None]   # header, getter, width, number_format


def _link(label: str, url: str | None):
    return ("__link__", label, url) if url else "—"


def _write_table(ws, title: str, subtitle: str, cols: list[Col], rows: list[dict], banner: str | None,
                 start_row: int = 1, freeze_col: int = 3, cf: Callable | None = None,
                 comments: Callable[[dict, str], str | None] | None = None, zebra: bool = True,
                 styler: Callable | None = None) -> None:
    r0 = start_row
    ws.cell(r0, 1, title).font = TITLE_FONT
    ws.cell(r0 + 1, 1, subtitle).font = SUB_FONT
    if banner:
        ws.cell(r0 + 2, 1, banner).font = WARN_FONT
    hdr = r0 + 3
    for j, (name, _, width, _) in enumerate(cols, start=1):
        c = ws.cell(hdr, j, name)
        c.font, c.fill = HEADER_FONT, HEADER_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(j)].width = width
    ws.row_dimensions[hdr].height = 30
    for i, row in enumerate(rows, start=hdr + 1):
        for j, (name, getter, _, fmt) in enumerate(cols, start=1):
            try:
                val = getter(row)
            except Exception:  # noqa: BLE001 — a formatting error must never kill the workbook
                val = "UNKNOWN"
            if val is None or (isinstance(val, float) and pd.isna(val)):
                val = "UNKNOWN"
            cell = ws.cell(i, j)
            if isinstance(val, tuple) and val and val[0] == "__link__":
                cell.value, cell.hyperlink, cell.font = val[1], val[2], LINK_FONT
            else:
                cell.value = val
                cell.font = BODY_FONT
            if fmt and isinstance(val, (int, float)):
                cell.number_format = fmt
            col_w = ws.column_dimensions[get_column_letter(j)].width or 8.43
            cell.alignment = WRAP if isinstance(val, str) and len(val) > col_w * 0.95 else TOP
            cell.border = BORDER
            if comments:
                note = comments(row, name)
                if note:
                    cell.comment = Comment(note, "pipeline")
            if styler:
                styler(cell, row, name)
    last = hdr + max(len(rows), 1)
    ref = f"A{hdr}:{get_column_letter(len(cols))}{last}"
    ws.freeze_panes = ws.cell(hdr + 1, freeze_col)
    if rows:
        ws.auto_filter.ref = ref
        body = f"A{hdr + 1}:{get_column_letter(len(cols))}{last}"
        # data missing → grey italic; alternating rows for readability
        ws.conditional_formatting.add(body, CellIsRule(operator="equal", formula=['"UNKNOWN"'],
                                                       font=Font(italic=True, color="8A8A8A"),
                                                       fill=PatternFill("solid", fgColor="F1F1EF")))
        if cf:
            cf(ws, hdr, last, cols)
        if zebra:
            ws.conditional_formatting.add(body, FormulaRule(formula=["MOD(ROW(),2)=0"],
                                                            fill=PatternFill("solid", fgColor="F6F8FA")))
    else:
        ws.cell(hdr + 1, 1, "No listings in this category for the current run.").font = SUB_FONT
    ws.sheet_view.zoomScale = 100


def _col_letter(cols: list[Col], name: str) -> str | None:
    for j, c in enumerate(cols, start=1):
        if c[0] == name:
            return get_column_letter(j)
    return None


def _standard_cf(ws, hdr: int, last: int, cols: list[Col], top5: bool = False):
    first = hdr + 1
    ncols = get_column_letter(len(cols))
    if top5:
        rank_col, tier_col = _col_letter(cols, "Rank"), _col_letter(cols, "Tier")
        ws.conditional_formatting.add(
            f"A{first}:{ncols}{last}",
            FormulaRule(formula=[f'AND(${tier_col}{first}="TOP 10",${rank_col}{first}<=5)'],
                        fill=PatternFill("solid", fgColor="E3F1E6"), stopIfTrue=False))
    noise = _col_letter(cols, "Noise Risk")
    if noise:
        rng = f"{noise}{first}:{noise}{last}"
        for word, fill, font in (("LIKELY QUIET", "D8F0DD", "0B6B22"), ("POSSIBLY QUIET", "E6F4E8", "2D6B3A"),
                                 ("NOISE UNCERTAIN", "ECEBE6", "4A4944"), ("LIKELY NOISY", "FBD5D2", "9B1C13")):
            ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT({noise}{first},{len(word)})="{word}"'],
                                                           fill=PatternFill("solid", fgColor=fill),
                                                           font=Font(bold=True, color=font), stopIfTrue=True))
    for name in ("Estimated Total USD", "Monthly Rent USD"):
        col = _col_letter(cols, name)
        if col:
            ws.conditional_formatting.add(
                f"{col}{first}:{col}{last}",
                CellIsRule(operator="greaterThan", formula=["1000"], fill=PatternFill("solid", fgColor="FDE2C8"),
                           font=Font(bold=True, color="9A3412"), stopIfTrue=True))
    bc = _col_letter(cols, "Budget Class")
    if bc:
        rng = f"{bc}{first}:{bc}{last}"
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f'{bc}{first}="STRICT_ALL_IN"'],
                                                       fill=PatternFill("solid", fgColor="D8F0DD"),
                                                       font=Font(bold=True, color="0B6B22"), stopIfTrue=True))
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f'{bc}{first}="BASE_RENT_COMPLIANT"'],
                                                       fill=PatternFill("solid", fgColor="FFF1CC"),
                                                       font=Font(bold=True, color="7A5200"), stopIfTrue=True))
    fit = _col_letter(cols, "Fit Score")
    if fit:
        ws.conditional_formatting.add(f"{fit}{first}:{fit}{last}",
                                      CellIsRule(operator="greaterThanOrEqual", formula=["75"],
                                                 font=Font(bold=True, color="0B6B22")))


def _rent_comment(row: dict, col: str) -> str | None:
    if col == "Monthly Rent USD" and row.get("rent_usd_basis") == "CALCULATED":
        note = (f"Published as S/ {C.num(row.get('rent_pen')):,.0f}. Converted at 1 USD = S/ "
                f"{C.num(row.get('fx_rate_usd_pen')):.3f} (see METHODOLOGY).")
        portal_usd, diff = C.num(row.get("rent_usd_published")), C.num(row.get("rent_usd_published_diff_pct"))
        if portal_usd is not None and diff is not None:
            note += f" The portal also shows USD {portal_usd:,.0f} ({diff:+.1f}% vs this rate)."
        return note
    if col == "Monthly Rent USD" and row.get("rent_pen_basis") == "PUBLISHED" and row.get("rent_usd_basis") == "PUBLISHED":
        return f"Listing publishes both USD and S/ {C.num(row.get('rent_pen')):,.0f}."
    return None


EXEC_COLS: list[Col] = [
    ("Tier", lambda r: r["_tier"], 11, None),
    ("Rank", lambda r: r["_rank"], 6, "0"),
    ("Fit Score", lambda r: C.num(r.get("fit_score")), 7, "0.0"),
    ("Property", C.property_label, 34, None),
    ("Monthly Rent USD", lambda r: C.num(r.get("rent_usd")), 11, USD_FMT),
    ("Maintenance", C.maintenance_text, 16, None),
    ("Estimated Total USD", lambda r: C.num(r.get("estimated_total_monthly_usd")), 11, USD_FMT),
    ("Budget Class", lambda r: r.get("budget_class"), 16, None),
    ("Bedrooms", lambda r: int(C.num(r.get("bedrooms")) or 0), 8, "0"),
    ("Area m²", lambda r: C.num(r.get("area_m2")), 8, "0"),
    ("USD/m²", lambda r: C.num(r.get("rent_usd_per_m2")), 8, PPM2_FMT),
    ("Furnished", C.furnished_text, 10, None),
    ("Parking", C.parking_text, 9, None),
    ("Floor", C.floor_text, 7, None),
    ("Noise Risk", C.noise_text, 22, None),
    ("Quietness Score", C.quietness_value, 9, "0"),
    ("Location", C.location_text, 34, None),
    ("Contract", C.contract_text, 12, None),
    ("Deposit", C.deposit_text, 14, None),
    ("Main Advantage", lambda r: r.get("main_advantage"), 40, None),
    ("Main Drawback", lambda r: r.get("main_drawback"), 36, None),
    ("Foreign Tenant", lambda r: r.get("foreign_tenant_friendliness") or "UNKNOWN", 14, None),
    ("Source", C.sources_text, 14, None),
    ("View Listing", lambda r: _link("View listing", r.get("source_url")), 12, None),
    ("WhatsApp", lambda r: _link("Open WhatsApp", C.whatsapp_url(r)), 14, None),
    ("Map", lambda r: _link("Open map", C.map_url(r)), 10, None),
    ("Active status", lambda r: r.get("active_status") or "UNKNOWN", 16, None),
    ("QA Status", lambda r: r.get("qa_status") or "NOT_CHECKED", 18, None),
]

DETAIL_COLS: list[Col] = [
    ("Rank", lambda r: r["_rank"], 6, "0"),
    ("Fit Score", lambda r: C.num(r.get("fit_score")), 7, "0.0"),
    ("Property", C.property_label, 34, None),
    ("Category", lambda r: r.get("category"), 11, None),
    ("Why in this sheet", lambda r: r.get("_why"), 30, None),
    ("Monthly Rent USD", lambda r: C.num(r.get("rent_usd")), 11, USD_FMT),
    ("Rent as published", lambda r: f"{r.get('rent_currency')} {C.num(r.get('rent_original')):,.0f}", 13, None),
    ("Rent USD basis", lambda r: r.get("rent_usd_basis"), 11, None),
    ("Rent PEN", lambda r: C.num(r.get("rent_pen")), 10, '"S/" #,##0'),
    ("Maintenance", C.maintenance_text, 16, None),
    ("Maintenance basis", lambda r: r.get("maintenance_pen_basis") or r.get("maintenance_usd_basis"), 11, None),
    ("Estimated Total USD", lambda r: C.num(r.get("estimated_total_monthly_usd")), 11, USD_FMT),
    ("Budget Class", lambda r: r.get("budget_class"), 16, None),
    ("Budget note", lambda r: r.get("budget_note"), 26, None),
    ("Bedrooms", lambda r: int(C.num(r.get("bedrooms")) or 0), 8, "0"),
    ("Bathrooms", lambda r: C.num(r.get("bathrooms")), 8, "0.#"),
    ("Area m²", lambda r: C.num(r.get("area_m2")), 8, "0"),
    ("Area basis", lambda r: r.get("area_basis"), 8, None),
    ("Space band", lambda r: r.get("space_band"), 12, None),
    ("USD/m²", lambda r: C.num(r.get("rent_usd_per_m2")), 8, PPM2_FMT),
    ("Value band", lambda r: r.get("value_band"), 18, None),
    ("Furnished", C.furnished_text, 10, None),
    ("Parking", C.parking_text, 9, None),
    ("Floor", C.floor_text, 7, None),
    ("Balcony/Terrace", lambda r: "Yes" if r.get("balcony") is True or r.get("terrace") is True else "UNKNOWN", 9, None),
    ("Laundry", lambda r: C.yn(r.get("laundry")), 8, None),
    ("Elevator", lambda r: C.yn(r.get("elevator")), 8, None),
    ("Security", lambda r: C.yn(r.get("security_24h")), 8, None),
    ("Pets", lambda r: C.yn(r.get("pets_allowed")), 7, None),
    ("Interior view", lambda r: C.yn(r.get("interior_view")), 8, None),
    ("Noise Risk", C.noise_text, 22, None),
    ("Quietness Score", C.quietness_value, 9, "0"),
    ("Quietness evidence", lambda r: r.get("quietness_reason"), 50, None),
    ("Nearest major road (m)", lambda r: C.num(r.get("dist_major_road_m")), 10, "0"),
    ("Nearest nightclub (m)", lambda r: C.num(r.get("dist_nightclub_m")), 10, "0"),
    ("Bars ≤150 m", lambda r: C.num(r.get("bars_within_m")), 8, "0"),
    ("Daily needs nearby", lambda r: r.get("location_notes") or "UNKNOWN", 30, None),
    ("Contract", C.contract_text, 12, None),
    ("Deposit", C.deposit_text, 14, None),
    ("Utilities included", C.utilities_text, 14, None),
    ("Published", lambda r: r.get("publication_date") or "UNKNOWN", 11, None),
    ("Active status", lambda r: r.get("active_status"), 14, None),
    ("Main Advantage", lambda r: r.get("main_advantage"), 40, None),
    ("Main Drawback", lambda r: r.get("main_drawback"), 36, None),
    ("Red flags", lambda r: r.get("red_flags") or "", 40, None),
    ("Foreign Tenant", lambda r: r.get("foreign_tenant_friendliness") or "UNKNOWN", 14, None),
    ("Foreign-tenant evidence", lambda r: r.get("foreign_tenant_evidence") or "—", 30, None),
    ("Agency / agent", lambda r: r.get("agency_name") or r.get("agent_name") or "UNKNOWN", 20, None),
    ("Phone", lambda r: r.get("phone") or "UNKNOWN", 14, None),
    ("Source", C.sources_text, 14, None),
    ("# Sources", lambda r: C.num(r.get("number_of_sources")), 7, "0"),
    ("Duplicate group", lambda r: r.get("duplicate_group_id"), 9, None),
    ("Duplicate conflicts", lambda r: r.get("dup_inconsistencies") or "", 24, None),
    ("View Listing", lambda r: _link("View listing", r.get("source_url")), 12, None),
    ("WhatsApp", lambda r: _link("Open WhatsApp", C.whatsapp_url(r)), 14, None),
    ("Map", lambda r: _link("Open map", C.map_url(r)), 10, None),
    ("Location source", lambda r: f"{r.get('coord_source') or 'none'} / {r.get('coord_precision') or '-'}", 16, None),
    ("QA Status", lambda r: r.get("qa_status") or "NOT_CHECKED", 18, None),
    ("QA notes", lambda r: r.get("qa_notes") or "", 30, None),
]


def _records(df: pd.DataFrame, why: Callable[[dict], str] | None = None) -> list[dict]:
    rows = []
    for k, r in enumerate(C.records(df), start=1):
        r["_rank"] = k
        r["_why"] = why(r) if why else ""
        rows.append(r)
    return rows


def _total_cell(r: dict):
    total = C.num(r.get("estimated_total_monthly_usd"))
    if r.get("maintenance_included_in_rent") is True:
        return C.num(r.get("rent_usd"))
    return total if total is not None else "Unknown (maint. n/p)"


def _contact_cell(r: dict):
    wa = C.whatsapp_url(r)
    if wa:
        return _link("WhatsApp", wa)
    if r.get("phone"):
        return f"Tel. {r['phone']}"
    return _link("Via portal", r.get("source_url"))


def _short(text, n: int) -> str:
    """Keep whole phrases ("; "-separated) up to about n characters — never cut a phrase in half."""
    parts = [p for p in str(text or "").split("; ") if p]
    out: list[str] = []
    for p in parts:
        if out and len("; ".join(out + [p])) > n:
            break
        out.append(p)
    return "; ".join(out)


CLIENT_COLS: list[Col] = [
    ("Rank", lambda r: r["_rank"], 5, "0"),
    ("Property", C.property_label, 25, None),
    ("Rent", lambda r: C.num(r.get("rent_usd")), 9, USD_FMT),
    ("Maintenance", C.maintenance_text, 13, None),
    ("Est. Total", _total_cell, 12, USD_FMT),
    ("Bedrooms", lambda r: int(C.num(r.get("bedrooms")) or 0), 9, "0"),
    ("Area m²", lambda r: C.num(r.get("area_m2")) or "Unknown", 8, "0"),
    ("Furnished", C.furnished_text, 11, None),
    ("Noise", C.noise_text, 15, None),
    ("Why It Stands Out", lambda r: _short(r.get("main_advantage"), 120), 30, None),
    ("Main Drawback", lambda r: _short(r.get("main_drawback"), 100), 26, None),
    ("Listing", lambda r: _link("Open", r.get("source_url")), 7, None),
    ("WhatsApp / Contact", _contact_cell, 14, None),
    ("Map", lambda r: _link("Map", C.map_url(r)), 6, None),
]
GREEN_FILL, AMBER_FILL = PatternFill("solid", fgColor="D8F0DD"), PatternFill("solid", fgColor="FFE8B3")


def _client_styler(cell, row: dict, col: str) -> None:
    if col == "Est. Total":
        strict = row.get("budget_class") == "STRICT_ALL_IN"
        cell.fill = GREEN_FILL if strict else AMBER_FILL
        cell.font = Font(name="Calibri", size=10, bold=True, color="0B6B22" if strict else "8A4B00")
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    elif col == "Noise":
        label = str(row.get("noise_label"))
        colors = {"LIKELY QUIET": ("D8F0DD", "0B6B22"), "POSSIBLY QUIET": ("E6F4E8", "2D6B3A"),
                  "NOISE UNCERTAIN": ("ECEBE6", "4A4944"), "LIKELY NOISY": ("FBD5D2", "9B1C13")}
        if label in colors:
            cell.fill = PatternFill("solid", fgColor=colors[label][0])
            cell.font = Font(name="Calibri", size=10, bold=True, color=colors[label][1])
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    elif col in ("Property", "Maintenance", "Why It Stands Out", "Main Drawback", "WhatsApp / Contact", "Furnished"):
        cell.alignment = Alignment(wrap_text=True, vertical="top")


def build_workbook(path: Path, ranked: pd.DataFrame, meta: dict, audit_rows: list[dict],
                   methodology: list[tuple[str, str]], banner: str | None = None) -> Path:
    wb = Workbook()
    stamp = f"Generated {meta['generated_at']} · FX 1 USD = S/ {meta['fx_rate']:.3f} ({meta['fx_source_short']})"
    primary = ranked[ranked["category"] == "PRIMARY"].sort_values("rank_in_category")
    top_n, alt_n = meta["top_n"], meta["alternatives_max"]

    # 0. CLIENT_TOP_PICKS — the simple view for the client
    ws = wb.active
    ws.title = "CLIENT_TOP_PICKS"
    picks = []
    for k, r in enumerate(C.records(primary.head(top_n)), start=1):
        r["_rank"] = k
        picks.append(r)
    _write_table(ws, "Top 10 apartments to contact first — Miraflores, rent ≤ USD 1,000",
                 "Est. Total: green = rent + maintenance ≤ USD 1,000 · amber = rent within budget but the total is above "
                 "USD 1,000 or maintenance is not published. Noise is an estimate — visit at rush hour and at night. "
                 f"{stamp}.", CLIENT_COLS, picks, banner, freeze_col=3, zebra=False, styler=_client_styler,
                 comments=lambda r, c: (r.get("budget_note") if c == "Est. Total" else _rent_comment(r, "Monthly Rent USD")
                                        if c == "Rent" else None))
    ws.sheet_view.zoomScale = 90
    ws.row_dimensions[2].height = 28
    ws.cell(2, 1).alignment = Alignment(wrap_text=False, vertical="top")

    # 1. EXECUTIVE_SHORTLIST
    ws = wb.create_sheet("EXECUTIVE_SHORTLIST")
    exec_rows = []
    for k, r in enumerate(C.records(primary.head(top_n + alt_n)), start=1):
        r["_tier"] = "TOP 10" if k <= top_n else "ALTERNATIVE"
        r["_rank"] = k
        exec_rows.append(r)
    _write_table(ws, "Miraflores rental shortlist — Top 10 and strong alternatives",
                 f"{stamp} · Budget-compliant (rent ≤ USD 1,000), 1–2 bedrooms, inside Miraflores. "
                 f"Availability as checked at the QA timestamp; not guaranteed.",
                 EXEC_COLS, exec_rows, banner, freeze_col=5,
                 cf=lambda ws_, h, l, c: _standard_cf(ws_, h, l, c, top5=True), comments=_rent_comment)

    # 2–4. detail sheets
    for name, cat, title, why in (
        ("ALL_MATCHES", "PRIMARY", "All budget-compliant Miraflores matches", None),
        ("STRETCH_NEGOTIABLE", "STRETCH", "Stretch / negotiable (rent USD 1,001–1,100) — kept separate",
         lambda r: f"Rent USD {C.num(r.get('rent_usd')):,.0f}: above the USD 1,000 target; only worth it if negotiable"),
        ("NEAR_MISSES", ("BORDERLINE", "NEAR_MISS"), "Near misses and borderline cases — never budget-compliant",
         lambda r: (f"BORDERLINE: {r.get('budget_note')}" if r.get("category") == "BORDERLINE"
                    else r.get("exclusion_reason") or "")),
    ):
        cats = cat if isinstance(cat, tuple) else (cat,)
        sub = ranked[ranked["category"].isin(cats)].sort_values(["category", "rank_in_category"])
        ws = wb.create_sheet(name)
        _write_table(ws, title, stamp, DETAIL_COLS, _records(sub, why), banner, freeze_col=4,
                     cf=lambda ws_, h, l, c: _standard_cf(ws_, h, l, c), comments=_rent_comment)

    # 5. SOURCE_AUDIT
    ws = wb.create_sheet("SOURCE_AUDIT")
    audit_cols: list[Col] = [(k, (lambda key: lambda r: r.get(key, ""))(k), w, None) for k, w in (
        ("Source", 16), ("Status", 10), ("Method", 26), ("Target search URL", 40), ("Records collected", 10),
        ("Public access", 18), ("Pagination", 14), ("Detail pages", 14), ("Contact data", 16),
        ("Coordinates", 14), ("Publication date", 14), ("Maintenance fee", 14), ("Completeness", 14),
        ("Limitations / errors", 60))]
    _write_table(ws, "Source audit", stamp, audit_cols, audit_rows, banner, freeze_col=2)

    # 6. METHODOLOGY
    ws = wb.create_sheet("METHODOLOGY")
    ws.cell(1, 1, "Methodology, scoring and limitations").font = TITLE_FONT
    ws.cell(2, 1, stamp).font = SUB_FONT
    if banner:
        ws.cell(3, 1, banner).font = WARN_FONT
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 120
    for i, (k, v) in enumerate(methodology, start=5):
        a, b = ws.cell(i, 1, k), ws.cell(i, 2, v)
        a.font = Font(name="Calibri", size=10, bold=True, color=NAVY)
        b.font = BODY_FONT
        a.alignment = b.alignment = WRAP
        b.border = a.border = BORDER

    # 7. CONTACT_GUIDE
    ws = wb.create_sheet("CONTACT_GUIDE")
    ws.cell(1, 1, "Contact guide — nothing has been sent; contact is left to you").font = TITLE_FONT
    ws.cell(2, 1, stamp).font = SUB_FONT
    if banner:
        ws.cell(3, 1, banner).font = WARN_FONT
    ws.column_dimensions["A"].width = 24
    ws.cell(5, 1, "Message to send (Spanish)").font = Font(bold=True, color=NAVY)
    ws.cell(5, 1).alignment = WRAP
    ws.cell(5, 2, TEMPLATE_ES).alignment = WRAP
    ws.merge_cells("B5:F5")
    ws.row_dimensions[5].height = 300
    ws.cell(6, 1, "What it says (English)").font = Font(bold=True, color=NAVY)
    ws.cell(6, 1).alignment = WRAP
    ws.cell(6, 2, TEMPLATE_EN).alignment = WRAP
    ws.merge_cells("B6:F6")
    ws.row_dimensions[6].height = 300
    contact_cols: list[Col] = [
        ("Rank", lambda r: r["_rank"], 10, "0"),
        ("Property", C.property_label, 34, None),
        ("Who to contact", C.contact_text, 36, None),
        ("View Listing", lambda r: _link("View listing", r.get("source_url")), 12, None),
        ("WhatsApp", lambda r: _link("Open WhatsApp (pre-filled ES message)", C.whatsapp_url(r)), 30, None),
        ("Ask specifically", lambda r: _specific_questions(r), 60, None),
    ]
    rows = _records(primary.head(top_n))
    _write_table(ws, "Top 10 — who to contact and what to ask", "", contact_cols, rows, None, start_row=8, freeze_col=1)
    ws.freeze_panes = None

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def _specific_questions(r: dict) -> str:
    q = []
    flags = str(r.get("red_flags") or "")
    if "UNKNOWN_MAINTENANCE" in flags:
        q.append("exact maintenance fee")
    if "AREA_UNKNOWN" in flags:
        q.append("floor area (m²)")
    if r.get("furnished") not in (True, False):
        q.append("furnished or not")
    if r.get("minimum_contract_months") is None or pd.isna(r.get("minimum_contract_months")):
        q.append("minimum lease term")
    if r.get("deposit_months") is None or pd.isna(r.get("deposit_months")):
        q.append("deposit months")
    if r.get("noise_risk") in ("MEDIUM", "HIGH") or r.get("interior_view") is not True:
        q.append("does the bedroom face the street or the interior?")
    if "APPROXIMATE_LOCATION" in flags or "INCOMPLETE_ADDRESS" in flags:
        q.append("exact address/block")
    return "; ".join(q) or "availability and visit times"
